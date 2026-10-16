"""Safe, deterministic extraction and chunking for supported document files."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path


SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".pdf", ".docx", ".xlsx", ".pptx"}

# These limits protect the API from unexpectedly large or pathological office
# documents even when the upload byte limit is raised for a deployment.
MAX_WORKBOOK_SHEETS = 50
MAX_SHEET_ROWS = 5_000
MAX_SHEET_COLUMNS = 100
MAX_WORKBOOK_CELLS = 500_000
MAX_EXTRACTED_CHARS = 2_000_000


def _ensure_not_empty(text: str) -> str:
    text = re.sub(r"\r\n?", "\n", text).strip()
    if not text:
        raise ValueError("The uploaded file is empty or contains no extractable text.")
    if len(text) > MAX_EXTRACTED_CHARS:
        raise ValueError(f"The extracted document exceeds the {MAX_EXTRACTED_CHARS:,} character limit.")
    return text


def _extract_pdf(content: bytes) -> str:
    try:
        import fitz  # PyMuPDF

        document = fitz.open(stream=content, filetype="pdf")
    except Exception as error:
        raise ValueError("The PDF file is malformed or could not be read.") from error
    try:
        try:
            pages = []
            for page_number, page in enumerate(document, start=1):
                page_text = page.get_text("text").strip()
                if page_text:
                    pages.append(f"[PDF page {page_number}]\n{page_text}")
            return _ensure_not_empty("\n\n".join(pages))
        except ValueError:
            raise
        except Exception as error:
            raise ValueError("The PDF file is malformed or could not be read.") from error
    finally:
        document.close()


def _iter_docx_blocks(document):
    """Yield paragraphs and tables in document order (rather than grouped order)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parent = document.element.body
    for child in parent.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _extract_docx(content: bytes) -> str:
    try:
        from docx import Document
        from docx.text.paragraph import Paragraph

        document = Document(io.BytesIO(content))
        blocks = []
        section_number = 1
        table_number = 0
        paragraph_number = 0
        for block in _iter_docx_blocks(document):
            if isinstance(block, Paragraph):
                value = block.text.strip()
                # In OOXML a section boundary is represented by sectPr on a
                # paragraph's properties. Headings are content, not section
                # boundaries, so they must not change citation locations.
                paragraph_properties = getattr(block._p, "pPr", None)
                if value:
                    paragraph_number += 1
                    blocks.append(f"[DOCX section {section_number}, paragraph {paragraph_number}]\n{value}")
                if paragraph_properties is not None and getattr(paragraph_properties, "sectPr", None) is not None:
                    section_number += 1
            else:
                table_number += 1
                rows = []
                for row_number, row in enumerate(block.rows, start=1):
                    values = [cell.text.strip() for cell in row.cells]
                    if any(values):
                        rows.append(
                            f"[DOCX section {section_number}, table {table_number} row {row_number}]\n"
                            + " | ".join(values)
                        )
                blocks.extend(rows)
        return _ensure_not_empty("\n\n".join(blocks))
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("The DOCX file is malformed or could not be read.") from error


def _extract_xlsx(content: bytes) -> str:
    workbook = None
    try:
        from openpyxl import load_workbook

        # read_only bounds memory; data_only never evaluates formulas, and the
        # default keep_vba=False ensures macro projects are not carried along.
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        if len(workbook.worksheets) > MAX_WORKBOOK_SHEETS:
            raise ValueError(f"The workbook contains more than {MAX_WORKBOOK_SHEETS} sheets.")
        output = []
        cell_count = 0
        for sheet in workbook.worksheets:
            rows = []
            for row_number, row in enumerate(
                sheet.iter_rows(max_row=MAX_SHEET_ROWS, max_col=MAX_SHEET_COLUMNS), start=1
            ):
                cell_count += len(row)
                if cell_count > MAX_WORKBOOK_CELLS:
                    raise ValueError(f"The workbook exceeds the {MAX_WORKBOOK_CELLS:,} cell limit.")
                values = ["" if cell.value is None else str(cell.value).strip() for cell in row]
                if any(values):
                    rows.append(f"[XLSX sheet {sheet.title}, row {row_number}]\n" + " | ".join(values))
            output.extend(rows)
        return _ensure_not_empty("\n\n".join(output))
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("The XLSX file is malformed or could not be read.") from error
    finally:
        # Read-only workbooks expose close(); tolerate a failure during an
        # already-failed load without masking the useful 400 error.
        try:
            workbook.close()  # type: ignore[union-attr]
        except Exception:
            pass


def _extract_pptx(content: bytes) -> str:
    def shape_text(shape) -> list[str]:
        if getattr(shape, "has_table", False):
            return [
                " | ".join(cell.text.strip() for cell in row.cells)
                for row in shape.table.rows
                if any(cell.text.strip() for cell in row.cells)
            ]
        if hasattr(shape, "text") and shape.text.strip():
            return [shape.text.strip()]
        if hasattr(shape, "shapes"):
            values = []
            for child in shape.shapes:
                values.extend(shape_text(child))
            return values
        return []

    try:
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(content))
        slides = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            parts = []
            for shape in slide.shapes:
                parts.extend(shape_text(shape))
            if parts:
                slides.append(f"[PPTX slide {slide_number}]\n" + "\n".join(parts))
        return _ensure_not_empty("\n\n".join(slides))
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("The PPTX file is malformed or could not be read.") from error


def extract_text(filename: str, content: bytes) -> str:
    """Decode a supported upload and normalize it into searchable text.

    CSV rows are rendered as pipe-separated values so the extracted text remains
    readable in evidence while retaining every cell.
    """
    suffix = Path(filename).suffix.lower()
    if suffix in {".pdf", ".docx", ".xlsx", ".pptx"}:
        return {
            ".pdf": _extract_pdf,
            ".docx": _extract_docx,
            ".xlsx": _extract_xlsx,
            ".pptx": _extract_pptx,
        }[suffix](content)

    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Keep this exception type for the route's clear UTF-8-specific 400.
        raise
    if suffix == ".csv":
        try:
            rows = csv.reader(io.StringIO(text), strict=True)
            text = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
        except csv.Error as error:
            raise ValueError("The CSV file is malformed.") from error
    return _ensure_not_empty(text)


def chunk_text(text: str, chunk_size: int = 1200, overlap: int = 200) -> list[str]:
    """Split text into deterministic overlapping chunks without empty chunks."""
    if chunk_size <= 0:
        raise ValueError("chunk_size must be positive")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap must be between zero and chunk_size - 1")

    normalized = re.sub(r"\s+", " ", text).strip()
    if not normalized:
        return []
    chunks: list[str] = []
    start = 0
    step = chunk_size - overlap
    while start < len(normalized):
        end = min(start + chunk_size, len(normalized))
        if end < len(normalized):
            boundary = normalized.rfind(" ", start + chunk_size // 2, end)
            if boundary > start:
                end = boundary
        chunks.append(normalized[start:end].strip())
        if end >= len(normalized):
            break
        start = max(end - overlap, start + step)
    return chunks
